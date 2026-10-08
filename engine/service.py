from __future__ import annotations

import base64
import hashlib
import io
import logging
import math
import multiprocessing as mp
import os
import queue
import re
import tempfile
import time
import unicodedata
from datetime import date, datetime
from pathlib import Path
from typing import Any

import pdfplumber
import pytesseract
from pytesseract import TesseractError, TesseractNotFoundError
from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Inches, Pt, RGBColor
from fastapi import FastAPI, HTTPException
from pdf2image import convert_from_path
from pdf2image.exceptions import PDFPopplerTimeoutError
from pydantic import BaseModel, Field
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

app = FastAPI(title="Career Hub Engine", version="1.0.0")
OCR_PAGE_TIMEOUT_SECONDS = int(os.getenv("OCR_PAGE_TIMEOUT_SECONDS", "20"))
PDF_RENDER_TIMEOUT_SECONDS = int(os.getenv("PDF_RENDER_TIMEOUT_SECONDS", "60"))
DEFAULT_COUNTRY = os.getenv("CAREER_DEFAULT_COUNTRY", "Portugal").strip() or "Portugal"


class ExtractRequest(BaseModel):
    filename: str
    data: str


class SearchRequest(BaseModel):
    query: str = Field(min_length=2, max_length=180)
    location: str = Field(default=DEFAULT_COUNTRY, max_length=120)
    sites: list[str] = Field(default_factory=lambda: ["indeed", "linkedin", "glassdoor"])
    results_wanted: int = Field(default=20, ge=1, le=50)
    hours_old: int | None = Field(default=168, ge=1, le=2160)
    remote: bool = False
    country: str = Field(default=DEFAULT_COUNTRY, max_length=80)
    geography_surface: str = Field(default="", max_length=80)
    portal_deadline_ms: int = Field(ge=1000, le=120000)
    engine_deadline_ms: int = Field(ge=2000, le=180000)


class GenerateRequest(BaseModel):
    full_name: str
    headline: str = ""
    contact_lines: list[str] = Field(default_factory=list)
    sections: dict[str, Any] = Field(default_factory=dict)
    output_name: str = "curriculo-adaptado"


SECTION_ALIASES = {
    "resumo": "summary",
    "resumo profissional": "summary",
    "resumo de competencias": "skills",
    "resumo de competências": "skills",
    "objetivo": "summary",
    "perfil profissional": "summary",
    "competencias": "skills",
    "competências": "skills",
    "habilidades": "skills",
    "conhecimentos": "skills",
    "skills": "skills",
    "skillset": "skills",
    "qualificacoes": "skills",
    "qualificações": "skills",
    "experiencia": "experience",
    "experiência": "experience",
    "experiencia profissional": "experience",
    "experiência profissional": "experience",
    "historico profissional": "experience",
    "formacao": "education",
    "formação": "education",
    "formacao academica": "education",
    "formação acadêmica": "education",
    "educacao": "education",
    "educação": "education",
    "certificacoes": "certifications",
    "certificações": "certifications",
    "certificados": "certifications",
    "cursos": "certifications",
    "idiomas": "languages",
    "projetos": "projects",
    "highlights": "projects",
}


def decode_payload(data: str) -> bytes:
    encoded = data.split(",", 1)[1] if data.startswith("data:") and "," in data else data
    try:
        return base64.b64decode(encoded, validate=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="Arquivo em base64 invalido.") from exc


def normalized_heading(value: str) -> str:
    value = re.sub(r"[:|]+$", "", value.strip().lower())
    value = unicodedata.normalize("NFKD", value)
    return "".join(char for char in value if not unicodedata.combining(char))


NORMALIZED_ALIASES = {normalized_heading(key): value for key, value in SECTION_ALIASES.items()}


def split_sections(text: str) -> dict[str, str]:
    sections: dict[str, list[str]] = {"header": []}
    active = "header"
    for raw_line in text.splitlines():
        line = raw_line.strip()
        key = NORMALIZED_ALIASES.get(normalized_heading(line))
        is_short_heading = key and len(line) <= 48
        if is_short_heading:
            active = key
            sections.setdefault(active, [])
            continue
        sections.setdefault(active, []).append(raw_line.rstrip())
    return {
        key: "\n".join(lines).strip()
        for key, lines in sections.items()
        if "\n".join(lines).strip()
    }


def extraction_quality_issues(text: str) -> int:
    collapsed_words = len(re.findall(r"[^\W\d_]{28,}", text, re.UNICODE))
    glued_punctuation = len(re.findall(r"[,;!?](?=[^\W\d_])", text, re.UNICODE))
    return collapsed_words * 3 + glued_punctuation


def normalize_ocr_text(text: str) -> str:
    return re.sub(r"([\w.+-]+@)\s+(?=[\w-])", r"\1", text)


def extract_pdf(path: Path) -> tuple[str, str, int]:
    pages: list[str] = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            pages.append(page.extract_text(x_tolerance=2, y_tolerance=3) or "")
        page_count = len(pdf.pages)
    text = "\n\n".join(part.strip() for part in pages if part.strip()).strip()
    method = "pdf-text"
    native_quality_issues = extraction_quality_issues(text)
    if len(re.sub(r"\s+", "", text)) < 80 or native_quality_issues:
        images = convert_from_path(
            path,
            dpi=220,
            first_page=1,
            last_page=min(page_count, 15),
            timeout=PDF_RENDER_TIMEOUT_SECONDS,
        )
        ocr_pages = [
            pytesseract.image_to_string(image, lang="por+eng", timeout=OCR_PAGE_TIMEOUT_SECONDS)
            for image in images
        ]
        ocr_text = normalize_ocr_text(
            "\n\n".join(part.strip() for part in ocr_pages if part.strip()).strip()
        )
        if len(ocr_text) > len(text) or extraction_quality_issues(ocr_text) < native_quality_issues:
            text = ocr_text
            method = "pdf-ocr-quality-fallback" if native_quality_issues else "pdf-ocr"
    return text, method, page_count


def extract_docx(path: Path) -> tuple[str, str, int]:
    document = Document(path)
    lines: list[str] = []
    for section in document.sections:
        lines.extend(paragraph.text for paragraph in section.header.paragraphs if paragraph.text.strip())
    lines.extend(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if values:
                lines.append(" | ".join(values))
    for section in document.sections:
        lines.extend(paragraph.text for paragraph in section.footer.paragraphs if paragraph.text.strip())
    return "\n".join(lines).strip(), "docx-xml", max(1, len(document.sections))


def extract_contact_data(text: str) -> dict[str, dict[str, Any]]:
    """Extract only contact values explicitly present in the document."""
    value = str(text or "")
    lines = [re.sub(r"\s+", " ", line).strip() for line in value.splitlines() if line.strip()]
    email_match = re.search(r"(?<![\w.+-])([\w.+-]+@[\w-]+(?:\.[\w-]+)+)(?![\w.-])", value, re.I)
    phone_match = re.search(r"(?<!\d)(?:\+?55[\s.-]?)?(?:\(?\d{2}\)?[\s.-]?)?9?\d{4}[\s.-]?\d{4}(?!\d)", value)
    linkedin_match = re.search(r"https?://(?:www\.)?linkedin\.com/in/[\w%._~:/?#\[\]@!$&'()*+,;=-]+", value, re.I)
    github_match = re.search(r"https?://(?:www\.)?github\.com/[\w.-]+(?:/[\w.-]+)?", value, re.I)
    location_match = re.search(r"\b(?:Bras[ií]lia(?:\s*[,\-–]\s*(?:DF|Distrito Federal))?|Distrito Federal(?:\s*,\s*Brasil)?)\b", value, re.I)
    name = ""
    rejected = re.compile(r"@|https?://|linkedin|github|curr[ií]culo|resume|telefone|phone|email|e-mail|\d{4}", re.I)
    for line in lines[:8]:
        if 2 <= len(line.split()) <= 7 and len(line) <= 100 and not rejected.search(line):
            name = line
            break

    def field(raw: str, confidence: float) -> dict[str, Any]:
        return {"value": raw.strip(), "source": "RESUME_PDF", "confidence": confidence}

    output: dict[str, dict[str, Any]] = {}
    if name:
        output["candidateName"] = field(name, 0.72)
    if email_match:
        output["email"] = field(email_match.group(1).lower(), 0.99)
    if phone_match:
        output["phone"] = field(phone_match.group(0), 0.94)
    if location_match:
        output["location"] = field(location_match.group(0), 0.9)
    if linkedin_match:
        output["linkedin"] = field(linkedin_match.group(0), 0.99)
    if github_match:
        output["github"] = field(github_match.group(0), 0.99)
    return output


def extract_file(filename: str, payload: bytes) -> dict[str, Any]:
    suffix = Path(filename).suffix.lower()
    if suffix not in {".pdf", ".docx", ".txt"}:
        raise HTTPException(status_code=400, detail="Use PDF, DOCX ou TXT.")
    temp_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as stream:
            temp_path = Path(stream.name)
            stream.write(payload)
        path = temp_path
        if suffix == ".pdf":
            text, method, pages = extract_pdf(path)
        elif suffix == ".docx":
            text, method, pages = extract_docx(path)
        else:
            text = payload.decode("utf-8", errors="replace").strip()
            method, pages = "plain-text", 1
    finally:
        if temp_path:
            temp_path.unlink(missing_ok=True)
    if len(text) < 40:
        detail = "O PDF nao possui texto legivel, mesmo apos o OCR." if suffix == ".pdf" else "Nao foi possivel extrair texto suficiente do arquivo."
        raise HTTPException(status_code=422, detail=detail)
    return {
        "text": text,
        "sections": split_sections(text),
        "contact": extract_contact_data(text),
        "method": method,
        "pages": pages,
        "characters": len(text),
    }


def clean_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if hasattr(value, "item"):
        try:
            return clean_value(value.item())
        except Exception:
            pass
    return value


def clean_string(value: Any) -> str:
    return re.sub(r"\s+", " ", str(clean_value(value) or "").replace("\ufffd", "")).strip()


def compact_location(record: dict[str, Any]) -> str:
    values = [record.get("location"), record.get("city"), record.get("state"), record.get("country")]
    unique: list[str] = []
    for value in values:
        value = clean_string(value)
        if value and value.lower() not in {item.lower() for item in unique}:
            unique.append(value)
    return ", ".join(unique)


def salary_text(record: dict[str, Any]) -> str:
    minimum = clean_value(record.get("min_amount"))
    maximum = clean_value(record.get("max_amount"))
    if minimum is None and maximum is None:
        return ""
    currency = clean_string(record.get("currency"))
    interval = clean_string(record.get("interval"))
    values = [f"{float(value):,.0f}".replace(",", ".") for value in [minimum, maximum] if value is not None]
    return " - ".join(values) + (f" {currency}" if currency else "") + (f" / {interval}" if interval else "")


def normalize_job(record: dict[str, Any]) -> dict[str, Any]:
    url = clean_string(record.get("job_url_direct") or record.get("job_url"))
    source = clean_string(record.get("site") or "desconhecido").lower()
    external_id = clean_string(record.get("id"))
    stable = external_id or hashlib.sha1(f"{source}|{url}".encode("utf-8")).hexdigest()[:20]
    return {
        "externalId": f"{source}:{stable}",
        "source": source,
        "title": clean_string(record.get("title")),
        "company": clean_string(record.get("company")),
        "location": compact_location(record),
        "city": clean_string(record.get("city")),
        "state": clean_string(record.get("state")),
        "country": clean_string(record.get("country")),
        "workModel": "Remoto" if bool(clean_value(record.get("is_remote"))) else "",
        "isRemote": bool(clean_value(record.get("is_remote"))),
        "salary": salary_text(record),
        "salaryMin": clean_value(record.get("min_amount")),
        "salaryMax": clean_value(record.get("max_amount")),
        "salaryCurrency": clean_string(record.get("currency")),
        "salaryInterval": clean_string(record.get("interval")),
        "description": clean_string(record.get("description")),
        "url": url,
        "sourceUrl": clean_string(record.get("job_url")),
        "applicationUrl": clean_string(record.get("job_url_direct")) or url,
        "datePosted": clean_value(record.get("date_posted")),
        "jobType": clean_string(record.get("job_type")),
    }


def safe_text(value: Any) -> str:
    if isinstance(value, list):
        value = "\n".join(str(item).strip() for item in value if str(item).strip())
    return (
        str(value or "")
        .replace("\u2022", "-")
        .replace("\u25cf", "-")
        .replace("\u25aa", "-")
        .replace("\u2013", "-")
        .replace("\u2014", "-")
        .strip()
    )


def section_lines(value: Any) -> list[str]:
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]
    return [line.strip(" \t-*") for line in str(value or "").splitlines() if line.strip(" \t-*")]


def add_docx_heading(document: Document, title: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(title.upper())
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(13, 79, 73)


def build_docx(request: GenerateRequest) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.05

    title = document.add_paragraph()
    title.alignment = 1
    title.paragraph_format.space_after = Pt(2)
    title_run = title.add_run(request.full_name or "Curriculo")
    title_run.bold = True
    title_run.font.name = "Arial"
    title_run.font.size = Pt(18)
    title_run.font.color.rgb = RGBColor(24, 32, 34)

    if request.headline:
        headline = document.add_paragraph()
        headline.alignment = 1
        headline.paragraph_format.space_after = Pt(2)
        run = headline.add_run(request.headline)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(23, 107, 99)

    if request.contact_lines:
        contact = document.add_paragraph()
        contact.alignment = 1
        contact.paragraph_format.space_after = Pt(7)
        run = contact.add_run(" | ".join(item for item in request.contact_lines if item))
        run.font.name = "Arial"
        run.font.size = Pt(9)

    labels = [
        ("summary", "Resumo profissional", "prose"),
        ("skills", "Competencias", "inline"),
        ("experience", "Experiencia profissional", "prose"),
        ("education", "Formacao", "prose"),
        ("certifications", "Certificacoes e cursos", "bullets"),
        ("projects", "Projetos", "prose"),
        ("languages", "Idiomas", "inline"),
    ]
    for key, label, mode in labels:
        value = request.sections.get(key)
        if not safe_text(value):
            continue
        add_docx_heading(document, label)
        if mode == "inline":
            document.add_paragraph(" | ".join(section_lines(value)))
        elif mode == "bullets":
            for line in section_lines(value):
                paragraph = document.add_paragraph(style="List Bullet")
                paragraph.paragraph_format.space_after = Pt(1)
                paragraph.add_run(line)
        else:
            blocks = [part.strip() for part in re.split(r"\n\s*\n", safe_text(value)) if part.strip()]
            for block in blocks:
                paragraph = document.add_paragraph()
                for index, line in enumerate(block.splitlines()):
                    if index:
                        paragraph.add_run().add_break(WD_BREAK.LINE)
                    paragraph.add_run(line.strip())

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def xml_escape(value: str) -> str:
    return (
        value.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )


def build_pdf(request: GenerateRequest) -> bytes:
    buffer = io.BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=16 * mm,
        leftMargin=16 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
        title=f"Curriculo - {request.full_name}",
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ResumeTitle", parent=styles["Title"], fontName="Helvetica-Bold",
        fontSize=17, leading=20, alignment=TA_CENTER, textColor=HexColor("#182022"),
        spaceAfter=3,
    )
    headline_style = ParagraphStyle(
        "ResumeHeadline", parent=styles["Normal"], fontName="Helvetica-Bold",
        fontSize=10.5, leading=13, alignment=TA_CENTER, textColor=HexColor("#176b63"),
        spaceAfter=3,
    )
    contact_style = ParagraphStyle(
        "ResumeContact", parent=styles["Normal"], fontName="Helvetica",
        fontSize=8.5, leading=11, alignment=TA_CENTER, spaceAfter=7,
    )
    heading_style = ParagraphStyle(
        "ResumeHeading", parent=styles["Heading2"], fontName="Helvetica-Bold",
        fontSize=10.5, leading=13, textColor=HexColor("#0d4f49"),
        spaceBefore=7, spaceAfter=3,
    )
    body_style = ParagraphStyle(
        "ResumeBody", parent=styles["BodyText"], fontName="Helvetica",
        fontSize=9.3, leading=12, spaceAfter=3,
    )
    bullet_style = ParagraphStyle(
        "ResumeBullet", parent=body_style, leftIndent=10, firstLineIndent=-6,
        bulletIndent=0, spaceAfter=1,
    )

    story: list[Any] = [Paragraph(xml_escape(request.full_name or "Curriculo"), title_style)]
    if request.headline:
        story.append(Paragraph(xml_escape(request.headline), headline_style))
    if request.contact_lines:
        story.append(Paragraph(xml_escape(" | ".join(request.contact_lines)), contact_style))

    labels = [
        ("summary", "RESUMO PROFISSIONAL", "prose"),
        ("skills", "COMPETENCIAS", "inline"),
        ("experience", "EXPERIENCIA PROFISSIONAL", "prose"),
        ("education", "FORMACAO", "prose"),
        ("certifications", "CERTIFICACOES E CURSOS", "bullets"),
        ("projects", "PROJETOS", "prose"),
        ("languages", "IDIOMAS", "inline"),
    ]
    for key, label, mode in labels:
        value = request.sections.get(key)
        if not safe_text(value):
            continue
        story.append(Paragraph(label, heading_style))
        if mode == "inline":
            story.append(Paragraph(xml_escape(" | ".join(section_lines(value))), body_style))
        elif mode == "bullets":
            for line in section_lines(value):
                story.append(Paragraph(xml_escape(line), bullet_style, bulletText="-"))
        else:
            for block in [part.strip() for part in re.split(r"\n\s*\n", safe_text(value)) if part.strip()]:
                story.append(Paragraph(xml_escape(block).replace("\n", "<br/>"), body_style))
    document.build(story)
    return buffer.getvalue()


@app.get("/health")
def health() -> dict[str, Any]:
    return {"ok": True, "services": ["documents", "job-search", "ocr", "generation"]}


@app.post("/documents/extract")
def extract_document(request: ExtractRequest) -> dict[str, Any]:
    try:
        return {"ok": True, **extract_file(request.filename, decode_payload(request.data))}
    except (TesseractError, TesseractNotFoundError) as exc:
        raise HTTPException(status_code=502, detail="Falha ao executar o OCR do PDF.") from exc
    except (RuntimeError, PDFPopplerTimeoutError) as exc:
        if "timeout" in str(exc).lower():
            raise HTTPException(status_code=504, detail="O OCR excedeu o tempo limite.") from exc
        raise


@app.post("/documents/generate")
def generate_document(request: GenerateRequest) -> dict[str, Any]:
    docx_data = build_docx(request)
    pdf_data = build_pdf(request)
    return {
        "ok": True,
        "docx": base64.b64encode(docx_data).decode("ascii"),
        "pdf": base64.b64encode(pdf_data).decode("ascii"),
    }


def portal_report(payload: dict[str, Any], status: str, started: float, started_at: str, **extra: Any) -> dict[str, Any]:
    elapsed = round((time.monotonic() - started) * 1000)
    count = max(0, int(extra.pop("count", 0)))
    return {
        "source": payload["site"],
        "site": payload["site"],
        "attempted": extra.pop("attempted", True),
        "status": status,
        "count": count,
        "rawResults": count,
        "acceptedResults": count,
        "durationMs": elapsed,
        "elapsedMs": elapsed,
        "query": payload["query"],
        "location": payload["location"],
        "country": payload["country"],
        "remote": payload["remote"],
        "geographySurface": payload["geography_surface"],
        "startedAt": started_at,
        "completedAt": datetime.utcnow().isoformat() + "Z",
        **extra,
    }


def build_jobspy_kwargs(payload: dict[str, Any]) -> dict[str, Any]:
    site = payload["site"]
    kwargs: dict[str, Any] = {
        "site_name": [site],
        "search_term": payload["query"],
        "location": payload["location"],
        "country_indeed": payload["country"],
        "results_wanted": payload["results_wanted"],
        "description_format": "markdown",
        "verbose": 0,
    }
    if site == "linkedin":
        kwargs["linkedin_fetch_description"] = False
    if payload.get("hours_old"):
        kwargs["hours_old"] = payload["hours_old"]
    if payload["remote"]:
        kwargs["is_remote"] = True
    if site == "google":
        kwargs["google_search_term"] = f"{payload['query']} vagas {payload['location']}"
    return kwargs


def execute_jobspy_portal(payload: dict[str, Any]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    from jobspy import scrape_jobs

    started = time.monotonic()
    started_at = datetime.utcnow().isoformat() + "Z"
    kwargs = build_jobspy_kwargs(payload)
    captured_errors: list[str] = []

    class PortalErrorHandler(logging.Handler):
        def emit(self, record: logging.LogRecord) -> None:
            if record.levelno >= logging.ERROR:
                captured_errors.append(clean_string(record.getMessage()))

    error_handler = PortalErrorHandler()
    portal_logger = logging.getLogger(f"JobSpy:{payload['site'].title()}")
    portal_logger.addHandler(error_handler)
    try:
        frame = scrape_jobs(**kwargs)
        raw_records = frame.to_dict(orient="records")
        records = [normalize_job(record) for record in raw_records]
        if payload["site"] == "linkedin":
            for record in records:
                record["provenance"] = {"linkedinSearch": {"remoteFilter": payload["remote"], "descriptionFetched": False}}
                if payload["remote"]:
                    record["isRemote"] = True
                    record["workModel"] = "Remoto"
        records = [record for record in records if record["title"] and record["company"]]
        if records:
            report = portal_report(
                payload,
                "SUCCESS",
                started,
                started_at,
                count=len(records),
                message=f"{len(records)} vaga(s) encontrada(s).",
            )
        elif captured_errors:
            detail = " ".join(captured_errors).lower()
            if "403" in detail or "forbidden" in detail or "blocked" in detail:
                status, code, message = "FAILED", "PORTAL_BLOCKED", "O portal bloqueou esta consulta automatizada."
            elif "location not parsed" in detail or "location" in detail and "400" in detail:
                status, code, message = "FAILED", "PORTAL_LOCATION_UNSUPPORTED", "O portal não reconheceu a localização solicitada."
            else:
                status, code, message = "FAILED", "PORTAL_FAILED", "A fonte falhou nesta tentativa."
            report = portal_report(payload, status, started, started_at, message=message, errorCode=code)
        else:
            report = portal_report(
                payload,
                "NO_RESULTS",
                started,
                started_at,
                message="Nenhuma vaga encontrada com estes filtros.",
            )
        report["rawResults"] = len(raw_records)
        report["normalizationLosses"] = len(raw_records) - len(records)
        return records, report
    except Exception as exc:
        lowered = clean_string(exc).lower()
        if "429" in lowered or "rate" in lowered or "too many" in lowered:
            status, code, message = "RATE_LIMITED", "PORTAL_RATE_LIMITED", "O portal limitou temporariamente as buscas."
        elif "timeout" in lowered or "timed out" in lowered:
            status, code, message = "TIMED_OUT", "PORTAL_TIMEOUT", "O portal demorou demais para responder."
        elif "403" in lowered or "forbidden" in lowered or "blocked" in lowered:
            status, code, message = "FAILED", "PORTAL_BLOCKED", "O portal bloqueou esta consulta automatizada."
        else:
            status, code, message = "FAILED", "PORTAL_FAILED", "A fonte falhou nesta tentativa."
        return [], portal_report(payload, status, started, started_at, message=message, errorCode=code)
    finally:
        portal_logger.removeHandler(error_handler)


def jobspy_portal_worker(payload: dict[str, Any], output: Any) -> None:
    if payload["site"] != "linkedin":
        output.put(execute_jobspy_portal(payload))
        return
    # Worker is a portal-isolated process. Count HTTP attempts, including retries.
    from urllib3.connectionpool import HTTPSConnectionPool
    original = HTTPSConnectionPool.urlopen
    requests_count = 0

    def counted(pool, *args, **kwargs):
        nonlocal requests_count
        if pool.host == "linkedin.com" or pool.host.endswith(".linkedin.com"):
            requests_count += 1
        return original(pool, *args, **kwargs)

    HTTPSConnectionPool.urlopen = counted
    try:
        records, report = execute_jobspy_portal(payload)
        report["externalRequests"] = requests_count
        output.put((records, report))
    finally:
        HTTPSConnectionPool.urlopen = original


@app.post("/jobs/linkedin/{job_id}/description")
def linkedin_description(job_id: str):
    from linkedin_enrichment import enrich_linkedin
    if not re.fullmatch(r"\d{1,20}", job_id):
        raise HTTPException(status_code=422, detail="Invalid LinkedIn job ID")
    return {"ok": True, **enrich_linkedin(job_id)}


def summarize_portal_outcomes(records: list[dict[str, Any]], reports: list[dict[str, Any]]) -> str:
    failed = any(item["status"] in {"FAILED", "TIMED_OUT", "RATE_LIMITED", "NOT_ATTEMPTED"} for item in reports)
    return "PARTIAL_SUCCESS" if records and failed else "SUCCESS" if records else "DEGRADED" if failed else "NO_RESULTS"


@app.post("/jobs/search")
def search_jobs(request: SearchRequest) -> dict[str, Any]:
    if request.portal_deadline_ms >= request.engine_deadline_ms:
        raise HTTPException(status_code=422, detail="portal_deadline_ms deve ser menor que engine_deadline_ms.")
    allowed = {"indeed", "linkedin", "glassdoor", "google"}
    sites = list(dict.fromkeys(site for site in request.sites if site in allowed)) or ["indeed"]
    concurrency = max(1, min(4, int(os.getenv("JOBSPY_PORTAL_CONCURRENCY", "2"))))
    operation_started = time.monotonic()
    operation_deadline = operation_started + request.engine_deadline_ms / 1000
    pending = list(sites)
    active: dict[str, dict[str, Any]] = {}
    records: list[dict[str, Any]] = []
    reports: list[dict[str, Any]] = []

    def payload_for(site: str) -> dict[str, Any]:
        return {
            "site": site,
            "query": request.query,
            "location": request.location or DEFAULT_COUNTRY,
            "country": request.country or DEFAULT_COUNTRY,
            "geography_surface": request.geography_surface,
            "remote": request.remote,
            "results_wanted": request.results_wanted,
            "hours_old": request.hours_old,
        }

    while pending or active:
        now = time.monotonic()
        while pending and len(active) < concurrency and now < operation_deadline:
            site = pending.pop(0)
            output = mp.Queue(maxsize=1)
            process = mp.Process(target=jobspy_portal_worker, args=(payload_for(site), output), daemon=True)
            process.start()
            active[site] = {
                "process": process,
                "output": output,
                "started": time.monotonic(),
                "started_at": datetime.utcnow().isoformat() + "Z",
            }
        progressed = False
        for site, state in list(active.items()):
            process, output = state["process"], state["output"]
            try:
                source_records, report = output.get_nowait()
                process.join(timeout=0.2)
                records.extend(source_records)
                reports.append(report)
                del active[site]
                progressed = True
                continue
            except queue.Empty:
                pass
            elapsed_ms = round((time.monotonic() - state["started"]) * 1000)
            if elapsed_ms >= request.portal_deadline_ms or time.monotonic() >= operation_deadline:
                process.terminate()
                process.join(timeout=1)
                payload = payload_for(site)
                reports.append(portal_report(payload, "TIMED_OUT", state["started"], state["started_at"], message="O portal excedeu o deadline isolado.", errorCode="PORTAL_TIMEOUT"))
                del active[site]
                progressed = True
            elif not process.is_alive():
                process.join(timeout=0.2)
                try:
                    source_records, report = output.get(timeout=0.2)
                    records.extend(source_records)
                    reports.append(report)
                except queue.Empty:
                    payload = payload_for(site)
                    reports.append(portal_report(payload, "FAILED", state["started"], state["started_at"], message="O processo isolado do portal terminou sem resposta.", errorCode="PORTAL_PROCESS_FAILED"))
                del active[site]
                progressed = True
        if time.monotonic() >= operation_deadline and pending:
            for site in pending:
                payload = payload_for(site)
                reports.append(portal_report(payload, "NOT_ATTEMPTED", operation_started, datetime.utcnow().isoformat() + "Z", attempted=False, message="O deadline da operação encerrou antes deste portal.", errorCode="ENGINE_DEADLINE_REACHED"))
            pending.clear()
        if not progressed:
            time.sleep(0.02)

    reports.sort(key=lambda item: sites.index(item["source"]))
    status = summarize_portal_outcomes(records, reports)
    return {
        "ok": True,
        "status": status,
        "jobs": records,
        "count": len(records),
        "sites": sites,
        "sourceReports": reports,
        "deadlines": {
            "portalDeadlineMs": request.portal_deadline_ms,
            "engineDeadlineMs": request.engine_deadline_ms,
        },
    }
