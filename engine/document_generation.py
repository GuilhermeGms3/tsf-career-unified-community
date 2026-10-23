import io
import re
from typing import Any

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate

from contracts import GenerateRequest


def safe_text(value: Any) -> str:
    if isinstance(value, list):
        value = "\n".join(str(item).strip() for item in value if str(item).strip())
    return str(value or "").replace("\u2022", "-").replace("\u25cf", "-").replace("\u25aa", "-").replace("\u2013", "-").replace("\u2014", "-").strip()


def section_lines(value: Any) -> list[str]:
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]
    return [line.strip(" \t-*") for line in str(value or "").splitlines() if line.strip(" \t-*")]


def add_docx_heading(document: Document, title: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(title.upper())
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(13, 79, 73)


def build_docx(request: GenerateRequest) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.05

    title = document.add_paragraph()
    title.alignment = 1
    title.paragraph_format.space_after = Pt(2)
    title_run = title.add_run(request.full_name or "Curriculo")
    title_run.bold = True
    title_run.font.name = "Arial"
    title_run.font.size = Pt(18)
    title_run.font.color.rgb = RGBColor(24, 32, 34)

    if request.headline:
        headline = document.add_paragraph()
        headline.alignment = 1
        headline.paragraph_format.space_after = Pt(2)
        run = headline.add_run(request.headline)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(23, 107, 99)

    if request.contact_lines:
        contact = document.add_paragraph()
        contact.alignment = 1
        contact.paragraph_format.space_after = Pt(7)
        run = contact.add_run(" | ".join(item for item in request.contact_lines if item))
        run.font.name = "Arial"
        run.font.size = Pt(9)

    labels = [("summary", "Resumo profissional", "prose"), ("skills", "Competencias", "inline"), ("experience", "Experiencia profissional", "prose"), ("education", "Formacao", "prose"), ("certifications", "Certificacoes e cursos", "bullets"), ("projects", "Projetos", "prose"), ("languages", "Idiomas", "inline")]
    for key, label, mode in labels:
        value = request.sections.get(key)
        if not safe_text(value):
            continue
        add_docx_heading(document, label)
        if mode == "inline":
            document.add_paragraph(" | ".join(section_lines(value)))
        elif mode == "bullets":
            for line in section_lines(value):
                paragraph = document.add_paragraph(style="List Bullet")
                paragraph.paragraph_format.space_after = Pt(1)
                paragraph.add_run(line)
        else:
            for block in [part.strip() for part in re.split(r"\n\s*\n", safe_text(value)) if part.strip()]:
                paragraph = document.add_paragraph()
                for index, line in enumerate(block.splitlines()):
                    if index:
                        paragraph.add_run().add_break(WD_BREAK.LINE)
                    paragraph.add_run(line.strip())
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def xml_escape(value: str) -> str:
    return value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")


def build_pdf(request: GenerateRequest) -> bytes:
    buffer = io.BytesIO()
    document = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=16 * mm, leftMargin=16 * mm, topMargin=14 * mm, bottomMargin=14 * mm, title=f"Curriculo - {request.full_name}")
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("ResumeTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=17, leading=20, alignment=TA_CENTER, textColor=HexColor("#182022"), spaceAfter=3)
    headline_style = ParagraphStyle("ResumeHeadline", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=10.5, leading=13, alignment=TA_CENTER, textColor=HexColor("#176b63"), spaceAfter=3)
    contact_style = ParagraphStyle("ResumeContact", parent=styles["Normal"], fontName="Helvetica", fontSize=8.5, leading=11, alignment=TA_CENTER, spaceAfter=7)
    heading_style = ParagraphStyle("ResumeHeading", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=10.5, leading=13, textColor=HexColor("#0d4f49"), spaceBefore=7, spaceAfter=3)
    body_style = ParagraphStyle("ResumeBody", parent=styles["BodyText"], fontName="Helvetica", fontSize=9.3, leading=12, spaceAfter=3)
    bullet_style = ParagraphStyle("ResumeBullet", parent=body_style, leftIndent=10, firstLineIndent=-6, bulletIndent=0, spaceAfter=1)

    story: list[Any] = [Paragraph(xml_escape(request.full_name or "Curriculo"), title_style)]
    if request.headline:
        story.append(Paragraph(xml_escape(request.headline), headline_style))
    if request.contact_lines:
        story.append(Paragraph(xml_escape(" | ".join(request.contact_lines)), contact_style))
    labels = [("summary", "RESUMO PROFISSIONAL", "prose"), ("skills", "COMPETENCIAS", "inline"), ("experience", "EXPERIENCIA PROFISSIONAL", "prose"), ("education", "FORMACAO", "prose"), ("certifications", "CERTIFICACOES E CURSOS", "bullets"), ("projects", "PROJETOS", "prose"), ("languages", "IDIOMAS", "inline")]
    for key, label, mode in labels:
        value = request.sections.get(key)
        if not safe_text(value):
            continue
        story.append(Paragraph(label, heading_style))
        if mode == "inline":
            story.append(Paragraph(xml_escape(" | ".join(section_lines(value))), body_style))
        elif mode == "bullets":
            for line in section_lines(value):
                story.append(Paragraph(xml_escape(line), bullet_style, bulletText="-"))
        else:
            for block in [part.strip() for part in re.split(r"\n\s*\n", safe_text(value)) if part.strip()]:
                story.append(Paragraph(xml_escape(block).replace("\n", "<br/>"), body_style))
    document.build(story)
    return buffer.getvalue()
