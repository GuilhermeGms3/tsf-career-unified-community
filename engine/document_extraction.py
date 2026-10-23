import base64
import os
import re
import tempfile
import unicodedata
from pathlib import Path
from typing import Any

import pdfplumber
import pytesseract
from docx import Document
from fastapi import HTTPException
from pdf2image import convert_from_path

OCR_PAGE_TIMEOUT_SECONDS = int(os.getenv("OCR_PAGE_TIMEOUT_SECONDS", "20"))
PDF_RENDER_TIMEOUT_SECONDS = int(os.getenv("PDF_RENDER_TIMEOUT_SECONDS", "60"))

SECTION_ALIASES = {
    "resumo": "summary", "resumo profissional": "summary",
    "resumo de competencias": "skills", "resumo de competências": "skills",
    "objetivo": "summary", "perfil profissional": "summary",
    "competencias": "skills", "competências": "skills", "habilidades": "skills",
    "conhecimentos": "skills", "skills": "skills", "skillset": "skills",
    "qualificacoes": "skills", "qualificações": "skills",
    "experiencia": "experience", "experiência": "experience",
    "experiencia profissional": "experience", "experiência profissional": "experience",
    "historico profissional": "experience", "formacao": "education",
    "formação": "education", "formacao academica": "education",
    "formação acadêmica": "education", "educacao": "education", "educação": "education",
    "certificacoes": "certifications", "certificações": "certifications",
    "certificados": "certifications", "cursos": "certifications",
    "idiomas": "languages", "projetos": "projects", "highlights": "projects",
}


def decode_payload(data: str) -> bytes:
    encoded = data.split(",", 1)[1] if data.startswith("data:") and "," in data else data
    try:
        return base64.b64decode(encoded, validate=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="Arquivo em base64 invalido.") from exc


def normalized_heading(value: str) -> str:
    value = re.sub(r"[:|]+$", "", value.strip().lower())
    value = unicodedata.normalize("NFKD", value)
    return "".join(char for char in value if not unicodedata.combining(char))


NORMALIZED_ALIASES = {normalized_heading(key): value for key, value in SECTION_ALIASES.items()}


def split_sections(text: str) -> dict[str, str]:
    sections: dict[str, list[str]] = {"header": []}
    active = "header"
    for raw_line in text.splitlines():
        line = raw_line.strip()
        key = NORMALIZED_ALIASES.get(normalized_heading(line))
        if key and len(line) <= 48:
            active = key
            sections.setdefault(active, [])
            continue
        sections.setdefault(active, []).append(raw_line.rstrip())
    return {key: "\n".join(lines).strip() for key, lines in sections.items() if "\n".join(lines).strip()}


def extraction_quality_issues(text: str) -> int:
    return len(re.findall(r"[^\W\d_]{28,}", text, re.UNICODE)) * 3 + len(re.findall(r"[,;!?](?=[^\W\d_])", text, re.UNICODE))


def normalize_ocr_text(text: str) -> str:
    return re.sub(r"([\w.+-]+@)\s+(?=[\w-])", r"\1", text)


def extract_pdf(path: Path) -> tuple[str, str, int]:
    pages: list[str] = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            pages.append(page.extract_text(x_tolerance=2, y_tolerance=3) or "")
        page_count = len(pdf.pages)
    text = "\n\n".join(part.strip() for part in pages if part.strip()).strip()
    method = "pdf-text"
    native_quality_issues = extraction_quality_issues(text)
    if len(re.sub(r"\s+", "", text)) < 80 or native_quality_issues:
        images = convert_from_path(path, dpi=220, first_page=1, last_page=min(page_count, 15), timeout=PDF_RENDER_TIMEOUT_SECONDS)
        ocr_pages = [pytesseract.image_to_string(image, lang="por+eng", timeout=OCR_PAGE_TIMEOUT_SECONDS) for image in images]
        ocr_text = normalize_ocr_text("\n\n".join(part.strip() for part in ocr_pages if part.strip()).strip())
        if len(ocr_text) > len(text) or extraction_quality_issues(ocr_text) < native_quality_issues:
            text = ocr_text
            method = "pdf-ocr-quality-fallback" if native_quality_issues else "pdf-ocr"
    return text, method, page_count


def extract_docx(path: Path) -> tuple[str, str, int]:
    document = Document(path)
    lines: list[str] = []
    for section in document.sections:
        lines.extend(paragraph.text for paragraph in section.header.paragraphs if paragraph.text.strip())
    lines.extend(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if values:
                lines.append(" | ".join(values))
    for section in document.sections:
        lines.extend(paragraph.text for paragraph in section.footer.paragraphs if paragraph.text.strip())
    return "\n".join(lines).strip(), "docx-xml", max(1, len(document.sections))


def extract_contact_data(text: str) -> dict[str, dict[str, Any]]:
    value = str(text or "")
    lines = [re.sub(r"\s+", " ", line).strip() for line in value.splitlines() if line.strip()]
    email_match = re.search(r"(?<![\w.+-])([\w.+-]+@[\w-]+(?:\.[\w-]+)+)(?![\w.-])", value, re.I)
    phone_match = re.search(r"(?<!\d)(?:\+?55[\s.-]?)?(?:\(?\d{2}\)?[\s.-]?)?9?\d{4}[\s.-]?\d{4}(?!\d)", value)
    linkedin_match = re.search(r"https?://(?:www\.)?linkedin\.com/in/[\w%._~:/?#\[\]@!$&'()*+,;=-]+", value, re.I)
    github_match = re.search(r"https?://(?:www\.)?github\.com/[\w.-]+(?:/[\w.-]+)?", value, re.I)
    location_match = re.search(r"\b(?:Bras[ií]lia(?:\s*[,\-–]\s*(?:DF|Distrito Federal))?|Distrito Federal(?:\s*,\s*Brasil)?)\b", value, re.I)
    rejected = re.compile(r"@|https?://|linkedin|github|curr[ií]culo|resume|telefone|phone|email|e-mail|\d{4}", re.I)
    name = next((line for line in lines[:8] if 2 <= len(line.split()) <= 7 and len(line) <= 100 and not rejected.search(line)), "")

    def field(raw: str, confidence: float) -> dict[str, Any]:
        return {"value": raw.strip(), "source": "RESUME_PDF", "confidence": confidence}

    output: dict[str, dict[str, Any]] = {}
    for key, match, confidence in [("candidateName", name, .72), ("email", email_match.group(1).lower() if email_match else "", .99), ("phone", phone_match.group(0) if phone_match else "", .94), ("location", location_match.group(0) if location_match else "", .9), ("linkedin", linkedin_match.group(0) if linkedin_match else "", .99), ("github", github_match.group(0) if github_match else "", .99)]:
        if match:
            output[key] = field(match, confidence)
    return output


def extract_file(filename: str, payload: bytes) -> dict[str, Any]:
    suffix = Path(filename).suffix.lower()
    if suffix not in {".pdf", ".docx", ".txt"}:
        raise HTTPException(status_code=400, detail="Use PDF, DOCX ou TXT.")
    temp_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as stream:
            temp_path = Path(stream.name)
            stream.write(payload)
        if suffix == ".pdf":
            text, method, pages = extract_pdf(temp_path)
        elif suffix == ".docx":
            text, method, pages = extract_docx(temp_path)
        else:
            text, method, pages = payload.decode("utf-8", errors="replace").strip(), "plain-text", 1
    finally:
        if temp_path:
            temp_path.unlink(missing_ok=True)
    if len(text) < 40:
        detail = "O PDF nao possui texto legivel, mesmo apos o OCR." if suffix == ".pdf" else "Nao foi possivel extrair texto suficiente do arquivo."
        raise HTTPException(status_code=422, detail=detail)
    return {"text": text, "sections": split_sections(text), "contact": extract_contact_data(text), "method": method, "pages": pages, "characters": len(text)}
